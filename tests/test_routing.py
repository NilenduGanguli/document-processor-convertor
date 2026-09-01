"""Tests for the Azure-only routing table in :mod:`dpc.pdfread` (SPEC-PMD-2 §2).

The contract under test is one sentence: **no input is read locally for its text unless a
route says so in the configuration**. PDFs and images always go to Azure Document
Intelligence; Office and text formats take the route their setting names; everything else is
a refusal a caller can read.

Nothing here opens a socket and nothing here needs a key. Azure is served by
:class:`httpx.MockTransport` behind ``pdfread.OcrClient`` (the seam named in §2.5(i)), the
external renderer is served by a two-line Python script this module writes into ``tmp_path``,
and every document is synthesised in-test. That is deliberate: the product path is
Azure-only, and the test path must never call Azure.
"""
from __future__ import annotations

import base64
import io
import logging
import shlex
import struct
import sys
import zipfile
from pathlib import Path
from typing import Any

import httpx
import openpyxl
import pymupdf
import pytest
from fastapi.testclient import TestClient

from dpc import db, pdfread, storage
from dpc.api import app
from dpc.config import Settings, get_settings
from dpc.ocr_client import OcrClient
from dpc.pdfread import NeedsRecognition, UnsupportedFormat, read_document

# ---------------------------------------------------------------------------
# Settings and the DI stub
# ---------------------------------------------------------------------------
LAYOUT_JOB: dict[str, Any] = {
    "status": "succeeded",
    "analyzeResult": {
        "apiVersion": "2024-11-30",
        "modelId": "prebuilt-layout",
        "content": "RECOGNISED TITLE",
        "pages": [
            {
                "pageNumber": 1,
                "width": 8.5,
                "height": 11.0,
                "unit": "inch",
                "angle": 0,
                "lines": [],
                "words": [],
            }
        ],
        "paragraphs": [
            {
                "content": "RECOGNISED TITLE",
                "role": "title",
                "boundingRegions": [
                    {"pageNumber": 1, "polygon": [1.0, 1.0, 4.0, 1.0, 4.0, 2.0, 1.0, 2.0]}
                ],
                "spans": [{"offset": 0, "length": 16}],
            }
        ],
    },
}


class DiStub:
    """A fake DI endpoint: records every submit, answers one poll with a canned job."""

    def __init__(self) -> None:
        self.submissions: list[tuple[bytes, str]] = []

    def handler(self, request: httpx.Request) -> httpx.Response:
        if request.method == "POST":
            self.submissions.append(
                (request.content, request.headers.get("content-type", ""))
            )
            return httpx.Response(202, headers={"Operation-Location": "https://di.example/op/1"})
        return httpx.Response(200, json=LAYOUT_JOB)

    @property
    def transport(self) -> httpx.MockTransport:
        return httpx.MockTransport(self.handler)


def make_settings(**overrides: Any) -> Settings:
    """Settings isolated from any .env file lying around the working directory."""
    return Settings(_env_file=None, **overrides)


def di_settings(**overrides: Any) -> Settings:
    values: dict[str, Any] = {
        "azure_di_endpoint": "https://di.example",
        "azure_di_key": "unit-test-key",
        "ocr_timeout_seconds": 5.0,
        "ocr_poll_interval_seconds": 0.0,
        "ocr_max_polls": 5,
    }
    values.update(overrides)
    return make_settings(**values)


def patch_ocr(monkeypatch: pytest.MonkeyPatch, stub: DiStub) -> None:
    """Route pdfread's internally-constructed OcrClient through the stub transport."""
    monkeypatch.setattr(
        pdfread, "OcrClient", lambda s: OcrClient(s, transport=stub.transport)
    )


def wired(monkeypatch: pytest.MonkeyPatch, **overrides: Any) -> tuple[DiStub, Settings]:
    stub = DiStub()
    patch_ocr(monkeypatch, stub)
    return stub, di_settings(**overrides)


# ---------------------------------------------------------------------------
# Document builders
# ---------------------------------------------------------------------------
def text_pdf(pages: int = 1) -> bytes:
    doc = pymupdf.open()
    for index in range(pages):
        page = doc.new_page(width=612, height=792)
        page.insert_text((72, 100), f"Page {index + 1} heading line")
    data: bytes = doc.tobytes()
    doc.close()
    return data


def png_bytes() -> bytes:
    pixmap = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 32, 32), False)
    pixmap.clear_with(200)
    return pixmap.tobytes("png")


def xlsx_bytes() -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Register"
    sheet.append(["Name", "Qty"])
    sheet.append(["widget", 3])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def bmp_bytes(width: int = 2, height: int = 2) -> bytes:
    """A real 24-bit BMP: BITMAPFILEHEADER + a 40-byte BITMAPINFOHEADER + padded pixel rows.

    Built by hand rather than sniffed from a prefix, because the prefix is the defect: ``BM``
    is two printable ASCII characters and the DIB header size behind it is the only thing
    that separates a bitmap from a sentence beginning "BMW".
    """
    row = width * 3
    padded = (row + 3) // 4 * 4
    pixels = b"".join(b"\x80" * row + b"\x00" * (padded - row) for _ in range(height))
    header = struct.pack("<IiiHHIIiiII", 40, width, height, 1, 24, 0, len(pixels), 0, 0, 0, 0)
    offset = 14 + len(header)
    return b"BM" + struct.pack("<IHHI", offset + len(pixels), 0, 0, offset) + header + pixels


HTML = b"<!doctype html><html><body><h1>Policy</h1><p>Body paragraph.</p></body></html>"
EML = (
    b"From: sender@example.test\r\n"
    b"To: recipient@example.test\r\n"
    b"Subject: Onboarding pack\r\n"
    b"Date: Mon, 1 Jan 2024 00:00:00 +0000\r\n"
    b"MIME-Version: 1.0\r\n"
    b"Content-Type: text/plain; charset=utf-8\r\n"
    b"\r\n"
    b"Account holder statement, first line.\r\n"
    b"Second line.\r\n"
)


# ---------------------------------------------------------------------------
# PDFs and images: always Azure
# ---------------------------------------------------------------------------
def test_pdf_always_goes_to_di_whole(monkeypatch: pytest.MonkeyPatch) -> None:
    stub, settings = wired(monkeypatch)
    data = text_pdf(2)
    _, provider = read_document(data, filename="form.pdf", settings=settings)
    assert provider == "azure_layout"
    assert stub.submissions == [(data, "application/pdf")]


@pytest.mark.parametrize(
    ("data", "expected"),
    [
        (png_bytes(), "image/png"),
        (b"\xff\xd8\xff\xe0" + b"\x00" * 32, "image/jpeg"),
        (b"II*\x00" + b"\x00" * 32, "image/tiff"),
        (bmp_bytes(), "image/bmp"),
        (b"\x00\x00\x00\x18ftypheic" + b"\x00" * 32, "image/heif"),
    ],
)
def test_every_accepted_image_type_goes_to_di(
    monkeypatch: pytest.MonkeyPatch, data: bytes, expected: str
) -> None:
    stub, settings = wired(monkeypatch)
    _, provider = read_document(data, filename=None, settings=settings)
    assert provider == "azure_layout"
    assert stub.submissions[0][1] == expected


@pytest.mark.parametrize(
    "opening",
    [
        b"BMW Finance Ltd - customer due diligence file\nAccount: 1234\n",
        b"BMO Harris statement of account\nHolder: A. Patel\n",
        b"BM Group Holdings - beneficial ownership\nShare: 51%\n",
    ],
)
def test_a_two_byte_bm_prefix_is_not_a_bitmap(
    monkeypatch: pytest.MonkeyPatch, opening: bytes
) -> None:
    """A KYC corpus is full of documents beginning "BM"; none of them is an image.

    The bare ``BM`` magic classified every one of these as ``image/bmp`` and, because the
    bytes outrank the filename, the ``.txt`` hint could not save them: they were submitted to
    DI as a bitmap and came back ``InvalidContent`` (502) or, against a lenient endpoint, as a
    valid-looking artifact with zero blocks. The whole document must survive instead.
    """
    stub, settings = wired(monkeypatch)
    body = opening * 20
    view, provider = read_document(body, filename="cdd_notes.txt", settings=settings)
    assert provider == "plain-text"
    assert stub.submissions == []
    assert view.blocks
    assert opening.decode().splitlines()[0] in view.text()


def test_a_real_bitmap_is_still_recognised_and_sent_to_di(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """The header check must not cost a real BMP its route — the DIB size is what decides."""
    stub, settings = wired(monkeypatch)
    _, provider = read_document(bmp_bytes(4, 4), filename=None, settings=settings)
    assert provider == "azure_layout"
    assert stub.submissions[0][1] == "image/bmp"


@pytest.mark.parametrize("dib_size", [12, 40, 52, 56, 64, 108, 124])
def test_every_real_dib_header_size_is_accepted(dib_size: int) -> None:
    """All six DIB header versions Microsoft shipped, plus BITMAPCOREHEADER."""
    data = b"BM" + b"\x00" * 12 + struct.pack("<I", dib_size) + b"\x00" * 8
    assert pdfread._classify(data, None) == ("image", "image/bmp")


def test_bytes_beat_the_filename(monkeypatch: pytest.MonkeyPatch) -> None:
    """A caller's name is a hint. PNG bytes named ``.pdf`` are submitted as ``image/png``."""
    stub, settings = wired(monkeypatch)
    read_document(png_bytes(), filename="statement.pdf", settings=settings)
    assert stub.submissions[0][1] == "image/png"


def test_corrupt_pdf_is_refused_before_any_network_call(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    stub, settings = wired(monkeypatch)
    with pytest.raises(ValueError, match="cannot open PDF"):
        read_document(b"%PDF-1.7 not really a pdf", filename="x.pdf", settings=settings)
    assert stub.submissions == []


def test_needs_recognition_names_configuration_and_counts_only() -> None:
    """KYC: the refusal a caller sees carries media type, page count and a variable name."""
    with pytest.raises(NeedsRecognition) as excinfo:
        read_document(text_pdf(2), filename="form.pdf", settings=make_settings())
    reason = excinfo.value.reason
    assert "application/pdf, 2 page(s)" in reason
    assert "DPC_AZURE_DI_ENDPOINT" in reason
    assert "heading line" not in reason  # no document text, ever


# ---------------------------------------------------------------------------
# Formats DI will not accept
# ---------------------------------------------------------------------------
@pytest.mark.parametrize(
    ("data", "filename"),
    [
        (b"GIF89a" + b"\x00" * 16, "scan.gif"),
        (b"RIFF\x00\x00\x00\x00WEBPVP8 ", "scan.webp"),
        (b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 16, "mail.msg"),
        (b"\x7fELF\x02\x01\x01" + b"\x00" * 16, "binary.bin"),
    ],
)
def test_formats_di_rejects_are_a_readable_refusal(data: bytes, filename: str) -> None:
    with pytest.raises(UnsupportedFormat) as excinfo:
        read_document(data, filename=filename, settings=di_settings())
    assert "image/jpeg" in excinfo.value.reason  # the message names what IS accepted


def test_gif_refusal_happens_before_the_submit(monkeypatch: pytest.MonkeyPatch) -> None:
    stub, settings = wired(monkeypatch)
    with pytest.raises(UnsupportedFormat):
        read_document(b"GIF89a" + b"\x00" * 16, filename=None, settings=settings)
    assert stub.submissions == []


# ---------------------------------------------------------------------------
# Office and HTML: DPC_OFFICE_ROUTE (§2.2)
# ---------------------------------------------------------------------------
def test_office_route_local_keeps_xlsx_tables(monkeypatch: pytest.MonkeyPatch) -> None:
    """The argued exception: DI returns no geometry for XLSX and no tables at all."""
    stub, settings = wired(monkeypatch)
    view, provider = read_document(xlsx_bytes(), filename="book.xlsx", settings=settings)
    assert provider == "openpyxl"
    assert view.tables  # the structure Azure would have dropped
    assert stub.submissions == []


def test_office_route_local_reads_html_locally(monkeypatch: pytest.MonkeyPatch) -> None:
    stub, settings = wired(monkeypatch)
    _, provider = read_document(HTML, filename="policy.html", settings=settings)
    assert provider == "htmlread"
    assert stub.submissions == []


def test_office_route_local_refuses_docx_with_the_two_remedies() -> None:
    """A DOCX has no local reader; the refusal names both routes that would convert it."""
    docx = _ooxml_package("word/document.xml")
    with pytest.raises(UnsupportedFormat) as excinfo:
        read_document(docx, filename="deed.docx", settings=di_settings())
    assert "DPC_OFFICE_ROUTE=render" in excinfo.value.reason
    assert "=azure" in excinfo.value.reason


def test_office_route_azure_submits_the_original_and_warns(
    monkeypatch: pytest.MonkeyPatch, caplog: pytest.LogCaptureFixture
) -> None:
    stub, settings = wired(monkeypatch, office_route="azure")
    data = xlsx_bytes()
    with caplog.at_level(logging.WARNING, logger="dpc.pdfread"):
        _, provider = read_document(data, filename="book.xlsx", settings=settings)
    assert provider == "azure_layout"
    assert stub.submissions[0][0] == data
    assert "spreadsheetml" in stub.submissions[0][1]
    assert any("azure_office_no_geometry" in record.message for record in caplog.records)


def _ooxml_package(member: str) -> bytes:
    """The smallest zip that classifies as a given OOXML kind — content, not filename."""
    import zipfile

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr(member, "<xml/>")
    return buffer.getvalue()


def _zip_package(members: dict[str, str]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        for name, body in members.items():
            archive.writestr(name, body)
    return buffer.getvalue()


def _content_types(main_part_media_type: str) -> str:
    return (
        '<?xml version="1.0"?><Types>'
        f'<Override PartName="/main" ContentType="{main_part_media_type}.main+xml"/>'
        "</Types>"
    )


def test_an_identifying_part_past_member_200_is_still_found(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """``namelist()[:200]`` was an arbitrary truncation of a writer-dependent ordering.

    Central-directory order is the writer's choice, so a media-heavy package whose
    identifying part sits past member 200 fell through to ``other`` -> 415. The cap bought
    nothing: ``any()`` over the full list already stops at the first match.
    """
    members = {"[Content_Types].xml": "<Types/>"}
    members.update({f"customXml/item{index}.xml": "<x/>" for index in range(250)})
    members["ppt/presentation.xml"] = "<p/>"
    assert pdfread._classify(_zip_package(members), None) == (
        "pptx",
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    )


def test_a_docx_with_an_embedded_workbook_is_not_read_as_a_spreadsheet(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """The package's own ``[Content_Types].xml`` outranks a prefix scan whose first hit wins.

    A deed with an embedded worksheet has both a ``word/`` and an ``xl/`` part, and the scan
    walked ``_OOXML`` in list order, so it was classified ``xlsx`` and handed to ``read_xlsx``.
    """
    docx_media = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    package = _zip_package(
        {
            "[Content_Types].xml": _content_types(docx_media),
            "xl/embeddings/Microsoft_Excel_Worksheet.xml": "<x/>",
            "word/document.xml": "<w/>",
        }
    )
    assert pdfread._classify(package, None) == ("docx", docx_media)
    # ...and it takes the DOCX route: local has no reader for it, and says so.
    with pytest.raises(UnsupportedFormat, match="DPC_OFFICE_ROUTE=render"):
        read_document(package, filename=None, settings=di_settings())


def test_an_ambiguous_package_with_no_content_types_refuses_rather_than_guessing() -> None:
    """Two prefixes and nothing authoritative: a readable 415 beats the wrong reader."""
    package = _zip_package(
        {"[Content_Types].xml": "<Types/>", "word/document.xml": "<w/>", "xl/book.xml": "<x/>"}
    )
    with pytest.raises(UnsupportedFormat) as excinfo:
        read_document(package, filename=None, settings=di_settings())
    assert "image/jpeg" in excinfo.value.reason  # the generic refusal, naming what IS accepted


def test_a_real_xlsx_is_still_classified_from_its_content_types(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Guard on the decisive path: openpyxl's own package must keep its route."""
    stub, settings = wired(monkeypatch)
    _, provider = read_document(xlsx_bytes(), filename=None, settings=settings)
    assert provider == "openpyxl"
    assert stub.submissions == []


def _fake_renderer(tmp_path: Path, *, exit_code: int = 0, pdfs: int = 1) -> str:
    """A DPC_RENDER_CMD that behaves like headless LibreOffice, without LibreOffice."""
    script = tmp_path / "renderer.py"
    script.write_text(
        "import sys, pymupdf\n"
        f"if {exit_code}:\n"
        f"    sys.exit({exit_code})\n"
        f"for n in range({pdfs}):\n"
        "    doc = pymupdf.open()\n"
        "    doc.new_page().insert_text((72, 100), 'rendered')\n"
        "    doc.save(sys.argv[2] + '/out%d.pdf' % n)\n"
        "    doc.close()\n",
        encoding="utf-8",
    )
    return f"{shlex.quote(sys.executable)} {shlex.quote(str(script))} {{input}} {{outdir}}"


def test_office_route_render_sends_the_rendered_pdf_and_stamps_the_renderer(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    stub, settings = wired(
        monkeypatch,
        office_route="render",
        render_cmd=_fake_renderer(tmp_path),
        render_version="libreoffice-24.2.5",
    )
    docx = _ooxml_package("word/document.xml")
    view, provider = read_document(docx, filename="deed.docx", settings=settings)
    assert provider == "azure_layout"
    body, content_type = stub.submissions[0]
    assert content_type == "application/pdf"
    assert body.startswith(b"%PDF-") and body != docx
    # The renderer version joins the determinism envelope, so it travels with the view.
    assert view.raw["renderer"] == "libreoffice-24.2.5"


def test_office_route_render_without_a_command_refuses() -> None:
    settings = di_settings(office_route="render", render_version="x")
    with pytest.raises(UnsupportedFormat, match="DPC_RENDER_CMD"):
        read_document(_ooxml_package("word/document.xml"), filename=None, settings=settings)


def test_office_route_render_without_a_version_stamp_refuses(tmp_path: Path) -> None:
    settings = di_settings(office_route="render", render_cmd=_fake_renderer(tmp_path))
    with pytest.raises(UnsupportedFormat, match="DPC_RENDER_VERSION"):
        read_document(_ooxml_package("word/document.xml"), filename=None, settings=settings)


def test_office_route_render_requires_the_placeholders(tmp_path: Path) -> None:
    settings = di_settings(
        office_route="render", render_cmd="soffice --headless", render_version="x"
    )
    with pytest.raises(ValueError, match=r"\{input\} and \{outdir\}"):
        read_document(_ooxml_package("word/document.xml"), filename=None, settings=settings)


def test_failed_render_reports_the_exit_status_and_no_renderer_output(
    tmp_path: Path,
) -> None:
    settings = di_settings(
        office_route="render",
        render_cmd=_fake_renderer(tmp_path, exit_code=3),
        render_version="x",
    )
    with pytest.raises(ValueError, match="exited 3"):
        read_document(_ooxml_package("word/document.xml"), filename=None, settings=settings)


def test_ambiguous_render_output_is_refused(tmp_path: Path) -> None:
    settings = di_settings(
        office_route="render",
        render_cmd=_fake_renderer(tmp_path, pdfs=2),
        render_version="x",
    )
    with pytest.raises(ValueError, match="produced 2 PDF"):
        read_document(_ooxml_package("word/document.xml"), filename=None, settings=settings)


# ---------------------------------------------------------------------------
# Text and mail: DPC_TEXT_ROUTE (§2.4)
# ---------------------------------------------------------------------------
def test_text_route_plain_claims_no_geometry(monkeypatch: pytest.MonkeyPatch) -> None:
    stub, settings = wired(monkeypatch)
    view, provider = read_document(
        b"Account holder: A\nBranch: B\n", filename="notes.txt", settings=settings
    )
    assert provider == "plain-text"
    assert stub.submissions == []
    assert view.blocks and all(block.bbox is None for block in view.blocks)
    assert view.raw["media_type"] == "text/plain"


def test_text_route_plain_needs_no_endpoint() -> None:
    """A .txt was convertible yesterday with no Azure account; it still is."""
    _, provider = read_document(b"one\ntwo\n", filename="a.txt", settings=make_settings())
    assert provider == "plain-text"


def test_email_takes_the_body_and_says_the_envelope_is_gone() -> None:
    view, provider = read_document(EML, filename="pack.eml", settings=make_settings())
    assert provider == "plain-text"
    assert view.raw["warning"] == "message envelope discarded"
    text = view.text()
    assert "Account holder statement, first line." in text
    assert "sender@example.test" not in text  # headers are discarded, not converted


LEDGER_CSV = b"Date:,From:,To:,Amount\n2024-01-01,ACME,BETA,10\n2024-01-02,GAMMA,DELTA,20\n"


@pytest.mark.parametrize("filename", ["ledger.csv", "ledger.txt", "ledger.md"])
def test_a_ledger_whose_header_row_looks_like_an_envelope_keeps_its_first_line(
    filename: str,
) -> None:
    """The mail sniff must never beat an explicit text extension (§2.4).

    ``Date:,From:,To:,Amount`` satisfied the old sniff — a colon in line 1 and two known
    header names anywhere in the block — so ``_message_body`` parsed the header row as an RFC
    5322 envelope and *deleted* it, leaving a "message envelope discarded" warning on a file
    that is not a message. Document text was dropped and nothing errored.
    """
    view, provider = read_document(LEDGER_CSV, filename=filename, settings=make_settings())
    assert provider == "plain-text"
    assert view.raw.get("warning") is None
    assert "Date:,From:,To:,Amount" in view.text()  # the header row survived
    assert "GAMMA" in view.text()


def test_an_unnamed_ledger_is_not_a_message_either() -> None:
    """With no filename the sniff is all there is, so the sniff itself must be structural.

    Every line of an RFC 5322 header block is a field or a folded continuation;
    ``2024-01-01,ACME,BETA,10`` is neither, so this is not mail whatever line 1 looks like.
    """
    view, provider = read_document(LEDGER_CSV, filename=None, settings=make_settings())
    assert provider == "plain-text"
    assert view.raw.get("warning") is None
    assert "Date:,From:,To:,Amount" in view.text()


def test_a_real_message_is_still_detected_from_its_bytes_alone() -> None:
    """The tightened sniff must not cost a genuine unnamed ``.eml`` its route."""
    view, provider = read_document(EML, filename=None, settings=make_settings())
    assert provider == "plain-text"
    assert view.raw["warning"] == "message envelope discarded"
    assert "Account holder statement, first line." in view.text()
    assert "sender@example.test" not in view.text()


def test_a_legitimate_non_utf8_message_converts_in_its_own_charset() -> None:
    """§2.4 says ``.eml`` is convertible; the UTF-8 gate belonged on the body, not the envelope.

    The gate ran on the raw message bytes, so any mail with an 8-bit body in its declared
    charset — French, German or Spanish KYC correspondence — was 415'd by a route the
    configuration says handles it. The stdlib parser decodes it correctly from the charset the
    message itself states.
    """
    eml = (
        "From: conformite@example.test\r\n"
        "To: kyc@example.test\r\n"
        "Subject: Dossier client\r\n"
        "MIME-Version: 1.0\r\n"
        "Content-Type: text/plain; charset=iso-8859-1\r\n"
        "\r\n"
        "Client: Herve Dupont, adresse verifiee.\r\n"
    ).encode("ascii").replace(b"Herve", b"Herv\xe9").replace(b"verifiee", b"v\xe9rifi\xe9e")
    view, provider = read_document(eml, filename="dossier.eml", settings=make_settings())
    assert provider == "plain-text"
    assert view.raw["warning"] == "message envelope discarded"
    assert "Hervé Dupont" in view.text()
    assert "vérifiée" in view.text()


def test_binary_in_a_decoded_message_body_is_still_refused() -> None:
    """Moving the gate must not open the route to binary: the body is checked, not skipped."""
    eml = (
        b"From: a@example.test\r\nTo: b@example.test\r\nSubject: x\r\n"
        b"MIME-Version: 1.0\r\nContent-Type: text/plain; charset=utf-8\r\n"
        b"Content-Transfer-Encoding: base64\r\n\r\n"
        + base64.b64encode(b"\x00\x01\x02 payload") + b"\r\n"
    )
    with pytest.raises(UnsupportedFormat):
        read_document(eml, filename="x.eml", settings=make_settings())


def test_an_html_fragment_named_html_takes_the_html_route(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """A CMS- or mail-exported fragment has no ``<html>`` tag; the extension is the evidence.

    Routed to plain text, its markup was emitted *as document text* — raw tags injected into
    the retrieval corpus, with htmlread's heading zoning and table extraction lost.
    """
    stub, settings = wired(monkeypatch)
    fragment = (
        b"<div><h1>Customer Due Diligence</h1>"
        b"<table><tr><td>Name</td><td>A. Patel</td></tr></table></div>\n"
    )
    view, provider = read_document(fragment, filename="policy.html", settings=settings)
    assert provider == "htmlread"
    assert stub.submissions == []
    assert "<h1>" not in view.text()
    assert "Customer Due Diligence" in view.text()


def test_text_route_refuse_names_both_remedies() -> None:
    settings = di_settings(text_route="refuse")
    with pytest.raises(UnsupportedFormat) as excinfo:
        read_document(b"plain report\n", filename="report.txt", settings=settings)
    assert "DPC_TEXT_ROUTE=plain" in excinfo.value.reason
    assert "=render" in excinfo.value.reason


def test_text_route_render_wraps_to_a_pdf_and_takes_the_pdf_route(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    stub, settings = wired(monkeypatch, text_route="render")
    report = "\n".join(f"ROW {index:04d}   {index * 7:>8}" for index in range(120))
    _, provider = read_document(report.encode(), filename="report.txt", settings=settings)
    assert provider == "azure_layout"
    body, content_type = stub.submissions[0]
    assert content_type == "application/pdf"
    with pymupdf.open(stream=body, filetype="pdf") as document:
        assert document.page_count == 3  # 120 rows at 54 rows/page
        assert "ROW 0000" in document[0].get_text()


def test_text_render_wrap_is_a_pure_function_of_the_text(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Determinism: the intermediate PDF carries no clock, so the same text wraps identically."""
    stub, settings = wired(monkeypatch, text_route="render")
    for _ in range(2):
        read_document(b"column one    column two\n", filename="r.txt", settings=settings)
    assert stub.submissions[0][0] == stub.submissions[1][0]


def test_text_render_over_max_pages_is_refused(monkeypatch: pytest.MonkeyPatch) -> None:
    stub, settings = wired(monkeypatch, text_route="render", max_pages=1)
    with pytest.raises(ValueError, match="over the 1-page limit"):
        read_document(b"line\n" * 200, filename="r.txt", settings=settings)
    assert stub.submissions == []


def test_the_form_feed_starts_a_new_page(monkeypatch: pytest.MonkeyPatch) -> None:
    """``\\f`` is the page break of the exact format this route exists for (§2.4).

    ``str.splitlines`` treats it as an ordinary line break, so a classic fixed-column
    report's own pagination was discarded and pages were re-cut every 54 rows.
    """
    stub, settings = wired(monkeypatch, text_route="render")
    report = "PAGE ONE HEADER\nrow one\fPAGE TWO HEADER\nrow two\fPAGE THREE HEADER\n"
    read_document(report.encode(), filename="report.txt", settings=settings)
    with pymupdf.open(stream=stub.submissions[0][0], filetype="pdf") as document:
        assert document.page_count == 3
        assert "PAGE ONE HEADER" in document[0].get_text()
        assert "PAGE TWO HEADER" in document[1].get_text()
        assert "PAGE TWO HEADER" not in document[0].get_text()
        assert "PAGE THREE HEADER" in document[2].get_text()


def test_a_form_feed_page_still_wraps_at_the_row_bound(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Honouring ``\\f`` must not remove the 54-row bound *within* a form-feed page."""
    stub, settings = wired(monkeypatch, text_route="render")
    long_page = "\n".join(f"ROW {index:04d}" for index in range(60))
    read_document((long_page + "\fsecond").encode(), filename="r.txt", settings=settings)
    with pymupdf.open(stream=stub.submissions[0][0], filetype="pdf") as document:
        assert document.page_count == 3  # 54 + 6 rows, then the form-feed page


@pytest.mark.parametrize(
    ("text", "expected_count"),
    [
        ("Иванов Иван", 10),
        ("Balance ₹1,00,000 — verified", 2),  # rupee sign and em dash
        ("余额", 2),
    ],
)
def test_render_route_refuses_glyphs_the_pinned_font_cannot_draw(
    monkeypatch: pytest.MonkeyPatch, text: str, expected_count: int
) -> None:
    """The pinned base-14 face has no coverage past Latin-1 and reports nothing when it misses.

    Rendered anyway, the document reaches Azure as a page of notdef dots, DI faithfully OCRs
    the dots, and a plausible PMD file is produced with the content gone and nothing warning
    about it. The refusal carries the COUNT and no document text (KYC: error bodies travel).
    """
    stub, settings = wired(monkeypatch, text_route="render")
    with pytest.raises(UnsupportedFormat) as excinfo:
        read_document(text.encode(), filename="r.txt", settings=settings)
    reason = excinfo.value.reason
    assert f"{expected_count} character(s)" in reason
    assert "DPC_TEXT_ROUTE=plain" in reason
    assert reason.isascii()  # not one character of the document is quoted back
    assert stub.submissions == []


def test_render_route_still_renders_everything_the_font_covers(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Latin-1 is real coverage, not an approximation: an accented name must survive."""
    stub, settings = wired(monkeypatch, text_route="render")
    read_document("Nom: Hervé Dupont".encode(), filename="r.txt", settings=settings)
    with pymupdf.open(stream=stub.submissions[0][0], filetype="pdf") as document:
        assert "Hervé Dupont" in document[0].get_text()


def test_text_render_refuses_for_want_of_an_endpoint_before_wrapping() -> None:
    """No endpoint means the wrapped PDF is worthless, so nothing is wrapped (§2.3).

    Pinned by the exception's identity: the ``max_pages`` ValueError below can only be raised
    by ``_wrap_text_to_pdf``, so seeing NeedsRecognition proves the wrap never ran.
    """
    settings = make_settings(text_route="render", max_pages=1)
    with pytest.raises(NeedsRecognition) as excinfo:
        read_document(b"line\n" * 200, filename="r.txt", settings=settings)
    assert "DPC_AZURE_DI_ENDPOINT" in excinfo.value.reason
    assert "text/plain" in excinfo.value.reason


def test_office_render_refuses_for_want_of_an_endpoint_before_the_subprocess() -> None:
    """Up to 120 s of LibreOffice per request to produce a 422 it could produce immediately.

    ``DPC_RENDER_CMD`` names an executable that does not exist, so reaching the subprocess
    raises ValueError; NeedsRecognition proves the check happened first.
    """
    settings = make_settings(
        office_route="render",
        render_cmd="/nonexistent/soffice --convert-to pdf --outdir {outdir} {input}",
        render_version="libreoffice-24.2.5",
    )
    with pytest.raises(NeedsRecognition) as excinfo:
        read_document(
            _ooxml_package("word/document.xml"), filename="deed.docx", settings=settings
        )
    assert "wordprocessingml" in excinfo.value.reason


def test_binary_bytes_never_reach_the_plain_text_route() -> None:
    """Valid UTF-8 is not enough: control characters mean binary, and binary is refused."""
    with pytest.raises(UnsupportedFormat):
        read_document(b"\x00\x01\x02payload", filename="report.txt", settings=di_settings())


# ---------------------------------------------------------------------------
# Config surface (§8)
# ---------------------------------------------------------------------------
def test_new_settings_carry_the_spec_defaults() -> None:
    settings = make_settings()
    assert settings.pmd_layout == "band"
    assert settings.pmd_rect_scale == "auto"
    assert settings.canvas_tab_snap is True
    assert settings.canvas_seg_rows == 20
    assert settings.canvas_seg_chars == 1400
    assert settings.canvas_row_y is False
    assert settings.canvas_emit_kv == "additive"
    assert settings.office_route == "local"
    assert settings.text_route == "plain"
    assert settings.render_cmd == ""
    assert settings.render_version == ""
    assert settings.allow_local_pdf_text is False
    assert settings.di_fixture_dir == ""


@pytest.mark.parametrize(
    ("field", "value"),
    [
        ("pmd_layout", "bands"),
        ("pmd_rect_scale", "milli"),
        ("office_route", "azur"),
        ("text_route", "plane"),
        ("canvas_emit_kv", "sometimes"),
        ("canvas_seg_rows", 0),
    ],
)
def test_nonsense_configuration_refuses_at_startup(field: str, value: Any) -> None:
    """Refuse-don't-guess: a typo must not silently convert half a corpus the other way."""
    with pytest.raises(ValueError):
        make_settings(**{field: value})


def test_enumerated_settings_are_case_and_space_insensitive() -> None:
    assert make_settings(office_route=" Azure ").office_route == "azure"


def test_retired_alnum_floor_is_inert_but_still_accepted() -> None:
    """An existing .env keeps validating; the floor decides nothing any more."""
    settings = make_settings(min_alnum_chars=999999)
    _, provider = read_document(b"still text\n", filename="a.txt", settings=settings)
    assert provider == "plain-text"


def test_get_settings_warns_that_the_retired_floor_is_inert(
    monkeypatch: pytest.MonkeyPatch, caplog: pytest.LogCaptureFixture
) -> None:
    """The deprecation WARNING lives in ``get_settings``, so only ``get_settings`` reaches it.

    Constructing ``Settings`` directly — which every other test here does — never runs that
    branch, so the operator-facing half of the retirement had no coverage at all.
    """
    monkeypatch.setenv("DPC_MIN_ALNUM_CHARS", "999")
    get_settings.cache_clear()
    try:
        with caplog.at_level(logging.WARNING, logger="dpc.config"):
            settings = get_settings()
        assert settings.min_alnum_chars == 999
        messages = [record.getMessage() for record in caplog.records]
        assert any("DPC_MIN_ALNUM_CHARS" in message for message in messages)
        assert any("effect=none" in message for message in messages)
    finally:
        get_settings.cache_clear()


def test_get_settings_is_quiet_when_the_retired_floor_is_not_set(
    monkeypatch: pytest.MonkeyPatch, caplog: pytest.LogCaptureFixture
) -> None:
    """A deployment that never set it must not be told about a setting it does not use."""
    monkeypatch.delenv("DPC_MIN_ALNUM_CHARS", raising=False)
    get_settings.cache_clear()
    try:
        with caplog.at_level(logging.WARNING, logger="dpc.config"):
            get_settings()
        assert not [
            record for record in caplog.records if "MIN_ALNUM" in record.getMessage()
        ]
    finally:
        get_settings.cache_clear()


# ---------------------------------------------------------------------------
# The HTTP layer: every §2.4 refusal is a 415 a client can read
#
# There was no HTTP-layer test for any refusal, which is exactly why the suite stayed green
# while the service answered 500 'internal' — a client error paging an on-call — and threw
# the carefully written remedy text away. These go through TestClient with the REAL
# dpc.pdfread, so the exception name the handler keys on is the real one.
# ---------------------------------------------------------------------------
@pytest.fixture()
def api_client(monkeypatch: pytest.MonkeyPatch) -> TestClient:
    """The app with storage and the database faked in memory. No lifespan, no sockets."""
    monkeypatch.setattr(storage, "put_markdown", lambda cid, text, settings=None: f"k/{cid}")
    monkeypatch.setattr(storage, "check", lambda settings=None: True)
    monkeypatch.setattr(db, "insert_conversion", lambda row, settings=None: None)
    monkeypatch.setattr(db, "init_schema_retrying", lambda *a, **k: None)
    monkeypatch.setattr(db, "check", lambda settings=None: True)
    # raise_server_exceptions=False because that is how a real client sees the app.
    return TestClient(app, raise_server_exceptions=False)


def post_document(client: TestClient, data: bytes, filename: str | None) -> httpx.Response:
    return client.post(
        "/api/v1/convert",
        json={
            "filename": filename,
            "content_base64": base64.b64encode(data).decode("ascii"),
        },
    )


@pytest.mark.parametrize(
    ("data", "filename", "expected_in_detail"),
    [
        # §2.4: a format DI does not accept and no route covers.
        (b"GIF89a" + b"\x00" * 16, "scan.gif", "image/gif"),
        (b"RIFF\x00\x00\x00\x00WEBPVP8 ", "scan.webp", "image/webp"),
        # Legacy OLE compound file: .doc/.xls/.ppt/.msg.
        (b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 16, "mail.msg", "x-ole-storage"),
        # §2.2: a DOCX under the default office_route=local has no local reader.
        (_ooxml_package("word/document.xml"), "deed.docx", "DPC_OFFICE_ROUTE=render"),
    ],
)
def test_every_refusal_family_is_a_415_over_http(
    api_client: TestClient, data: bytes, filename: str, expected_in_detail: str
) -> None:
    response = post_document(api_client, data, filename)
    assert response.status_code == 415
    body = response.json()
    assert body["error"] == "unsupported_media_type"
    assert expected_in_detail in body["detail"]  # the remedy survives to the client


def test_text_route_refuse_is_a_415_over_http_with_both_remedies(
    api_client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(
        "dpc.api.get_settings", lambda: make_settings(text_route="refuse")
    )
    response = post_document(api_client, b"plain report\n", "report.txt")
    assert response.status_code == 415
    assert "DPC_TEXT_ROUTE=plain" in response.json()["detail"]


def test_a_bm_document_is_converted_over_http_not_refused(api_client: TestClient) -> None:
    """The end-to-end shape of the BMP defect: a 'BMW…' file must convert, not 415/502."""
    response = post_document(
        api_client, b"BMW Finance Ltd - due diligence\nAccount: 1234\n" * 20, "cdd.txt"
    )
    assert response.status_code == 200
    body = response.json()
    assert body["provider"] == "plain-text"
    assert body["blocks"] > 0


def test_no_endpoint_is_still_a_422_over_http(api_client: TestClient) -> None:
    """The 415 mapping must not swallow the §2.3 refusal, which is a different answer."""
    response = post_document(api_client, text_pdf(1), "form.pdf")
    assert response.status_code == 422
    assert response.json()["error"] == "needs_ocr"


def test_a_refusal_body_carries_no_document_text(api_client: TestClient) -> None:
    """KYC: error bodies travel. A refusal names media types and variables, never content."""
    secret = b"Account holder Priya Raman, PAN ABCDE1234F"
    response = post_document(api_client, b"GIF89a" + secret, "scan.gif")
    assert response.status_code == 415
    assert "Priya" not in response.text and "ABCDE1234F" not in response.text


@pytest.mark.xfail(
    strict=True,
    reason=(
        "§2.4 handoff, owned by the lead: dpc/api.py must surface view.raw['warning'] in the "
        "convert response. pdfread sets it; nothing reads it. Flip this marker when wired."
    ),
)
def test_the_envelope_warning_reaches_the_caller(api_client: TestClient) -> None:
    response = post_document(api_client, EML, "pack.eml")
    assert response.status_code == 200
    assert response.json()["warning"] == "message envelope discarded"


@pytest.mark.xfail(
    strict=True,
    reason=(
        "§2.2 handoff, owned by the lead: dpc/emitter.py must publish view.raw['renderer'] as "
        "`renderer: <version>` in the front matter. pdfread stamps it; nothing reads it."
    ),
)
def test_the_renderer_version_reaches_the_front_matter(
    api_client: TestClient, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    stub = DiStub()
    patch_ocr(monkeypatch, stub)
    settings = di_settings(
        office_route="render",
        render_cmd=_fake_renderer(tmp_path),
        render_version="libreoffice-24.2.5",
    )
    monkeypatch.setattr("dpc.api.get_settings", lambda: settings)
    response = api_client.post(
        "/api/v1/convert",
        json={
            "filename": "deed.docx",
            "echo": True,
            "content_base64": base64.b64encode(
                _ooxml_package("word/document.xml")
            ).decode("ascii"),
        },
    )
    assert response.status_code == 200
    assert "renderer: libreoffice-24.2.5" in response.json()["markdown"]


def test_a_package_whose_content_types_part_is_unreadable_still_routes_by_prefix() -> None:
    """A damaged content-types part must reach the prefix-scan fallback, not fail the package.

    The read sat inside the outer ``except Exception: return None`` guarded only by
    ``KeyError``, so a corrupt deflate stream (BadZipFile) or a password-protected member
    (RuntimeError) discarded ``names`` as well — and an unnamed DOCX whose ``word/`` members
    were perfectly intact stopped being a DOCX and became an HTTP 415. The docstring promises
    the prefix scan "for packages with no usable content-types part"; unreadable is exactly
    that.
    """
    import io
    import zipfile

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types/>" * 50)
        archive.writestr("word/document.xml", "<w:document/>")
    raw = bytearray(buffer.getvalue())
    second = raw.find(b"PK\x03\x04", raw.find(b"PK\x03\x04") + 4)
    raw[second - 12] ^= 0xFF  # corrupt the first member's deflate stream

    kind = pdfread._ooxml_kind(bytes(raw))
    assert kind is not None, "an intact word/ package became unroutable"
    assert kind[0] == "docx"
